import asyncio
import datetime
import logging
import os
import uuid
import json
from abc import ABC, abstractmethod
from typing import Dict, Any, List, Optional
from decimal import Decimal
from io import BytesIO

import dotenv
from apscheduler.schedulers.blocking import BlockingScheduler
from apscheduler.triggers.interval import IntervalTrigger
from pymongo import AsyncMongoClient
import hvac

# Database connectors
import pymysql
import pymssql
import asyncpg

# Google Drive
import googleapiclient.discovery
from google.oauth2 import service_account
from docx import Document
from pptx import Presentation
from googleapiclient.http import MediaIoBaseDownload

# OpenSearch and vectorization
from opensearchpy._async.client import AsyncOpenSearch
import requests

# Настройка логирования
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

dotenv.load_dotenv()

# Конфигурация
SCHEDULER_INTERVAL_HOURS = 10
MONGO_DB_NAME = 'source-managment'
MONGO_COLLECTION_NAME = 'resources'
# VAULT_URL = os.getenv('VAULT_URL', 'http://localhost:8200')
# VAULT_TOKEN = os.getenv('VAULT_TOKEN', 'root')
VAULT_TOKEN = 'root'
VAULT_URL = 'http://localhost:8200'
# OPENSEARCH_HOST = os.getenv('OPENSEARCH_HOST', 'localhost:9200')
MONGO_URL = os.getenv('MONGO_URL', 'mongodb://localhost:27017')
OPENSEARCH_HOST = 'localhost:9200'


class DecimalEncoder(json.JSONEncoder):
    """JSON Encoder для обработки Decimal и datetime объектов"""

    def default(self, o):
        if isinstance(o, Decimal):
            return str(o)
        elif isinstance(o, datetime.datetime):
            return o.isoformat()
        return super().default(o)


class VectorizationService:
    """Сервис для векторизации данных"""

    def __init__(self, vault_client: hvac.Client):
        self.vault_client = vault_client
        self._vectorization_url = os.getenv('VECTORIZE_SERVICE_URL', 'http://localhost:8010/api/v1/vectorize_text')

    def _get_vectorization_url(self) -> str:
        """Получение URL сервиса векторизации"""
        if not self._vectorization_url:
            secrets = self.vault_client.secrets.kv.v2.read_secret_version(
                path='services/vectorization'
            )['data']['data']
            self._vectorization_url = secrets.get('vectorize_service_url')
        return self._vectorization_url

    def vectorize_data(self, content: str) -> List[float]:
        """Векторизация текстового контента"""
        try:
            logger.info("Starting vectorization")
            vectorization_url = self._get_vectorization_url()
            params = {"text": content}

            response = requests.post(vectorization_url, json=params, timeout=60)
            response.raise_for_status()

            vectorized_data = response.json().get("vectorized_text")
            logger.info("Vectorization completed successfully")
            return vectorized_data

        except Exception as e:
            logger.error(f"Vectorization failed: {e}")
            raise


class OpenSearchService:
    """Сервис для работы с OpenSearch"""

    def __init__(self, vault_client: hvac.Client):
        self.vault_client = vault_client
        self._client = None

    async def _get_client(self) -> AsyncOpenSearch:
        """Получение клиента OpenSearch"""
        if not self._client:
            secrets = self.vault_client.secrets.kv.v2.read_secret_version(
                path='search/opensearch'
            )['data']['data']

            self._client = AsyncOpenSearch(
                hosts=[secrets.get("opensearch_host", OPENSEARCH_HOST)],
                http_auth=(
                    secrets.get("opensearch_username"),
                    secrets.get("opensearch_password")
                ),
                use_ssl=True,
                verify_certs=False,
                timeout=60,
                max_retries=10,
                retry_on_timeout=True,
            )
        return self._client

    async def delete_by_resource_id(self, resource_id: str, index: str):
        """Удаление всех документов по resource_id"""
        try:
            client = await self._get_client()
            query = {
                "query": {
                    "match": {"resource_id": resource_id}
                }
            }
            await client.delete_by_query(index=index, body=query)
            logger.info(f"Deleted existing data for resource {resource_id}")
        except Exception as e:
            logger.error(f"Failed to delete data for resource {resource_id}: {e}")

    async def add_document(self, resource_id: str, index: str, vector: List[float],
                           content: Any, sub_id: str = None):
        """Добавление документа в OpenSearch"""
        try:
            client = await self._get_client()
            doc_id = sub_id or str(uuid.uuid4())

            # Обработка разных типов контента
            if isinstance(content, (dict, list)):
                try:
                    content_str = json.dumps(content, cls=DecimalEncoder, ensure_ascii=False)
                except (TypeError, ValueError):
                    content_str = str(content)
            else:
                content_str = str(content)

            document = {
                "resource_id": resource_id,
                "sub_resource_id": doc_id,
                "vector": vector,
                "content": content_str,
                "processing_timestamp": datetime.datetime.now().isoformat()
            }

            await client.update(
                index=index,
                id=doc_id,
                body={
                    "doc": document,
                    "doc_as_upsert": True,
                },
                timeout=60
            )

            logger.info(f"Successfully added document {doc_id} for resource {resource_id}")

        except Exception as e:
            logger.error(f"Failed to add document to OpenSearch: {e}")
            raise


# Стратегии обработки ресурсов
class ResourceProcessingStrategy(ABC):
    """Абстрактная стратегия обработки ресурсов"""

    def __init__(self, vault_client: hvac.Client, vectorization_service: VectorizationService,
                 opensearch_service: OpenSearchService):
        self.vault_client = vault_client
        self.vectorization_service = vectorization_service
        self.opensearch_service = opensearch_service

    @abstractmethod
    async def process_resource(self, resource: Dict[str, Any]):
        """Обработка ресурса"""
        pass

    def _get_secret(self, secret_path: str) -> Dict[str, Any]:
        """Получение секретов из Vault"""
        response = self.vault_client.secrets.kv.v2.read_secret_version(path=secret_path)
        return response['data']['data']


class MySQLProcessingStrategy(ResourceProcessingStrategy):
    """Стратегия обработки MySQL ресурсов"""

    async def process_resource(self, resource: Dict[str, Any]):
        logger.info(f"Processing MySQL resource: {resource.get('resource_id')}")

        secret = self._get_secret(resource['extra'].get("database_secret_path"))
        connection_params = secret.get("connection_params")

        connection = pymysql.connect(
            host=connection_params.get("host"),
            port=connection_params.get("port"),
            user=connection_params.get("user"),
            password=connection_params.get("password"),
            db=connection_params.get("db"),
            cursorclass=pymysql.cursors.DictCursor
        )

        try:
            # Удаляем существующие данные
            await self.opensearch_service.delete_by_resource_id(
                resource.get("resource_id"), "vectorized_knowledge_resource"
            )

            with connection.cursor() as cursor:
                cursor.execute(resource['extra'].get("query"))
                batch_size = resource['extra'].get("batch_size", 100)

                while True:
                    batch_data = cursor.fetchmany(batch_size)
                    if not batch_data:
                        break

                    for row in batch_data:
                        content_str = json.dumps(row, cls=DecimalEncoder, ensure_ascii=False)
                        vector = self.vectorization_service.vectorize_data(content_str)

                        await self.opensearch_service.add_document(
                            resource.get("resource_id"),
                            "vectorized_knowledge_resource",
                            vector,
                            row,
                            str(uuid.uuid4())
                        )

                    logger.info(f"Processed batch of {len(batch_data)} records")

        finally:
            connection.close()


class PostgreSQLProcessingStrategy(ResourceProcessingStrategy):
    """Стратегия обработки PostgreSQL ресурсов"""

    async def process_resource(self, resource: Dict[str, Any]):
        logger.info(f"Processing PostgreSQL resource: {resource.get('resource_id')}")

        secret = self._get_secret(resource['extra'].get("database_secret_path"))
        connection_params = secret.get("connection_params")

        connection = await asyncpg.connect(
            host=connection_params.get("host"),
            port=connection_params.get("port"),
            user=connection_params.get("user"),
            password=connection_params.get("password"),
            database=connection_params.get("database")
        )

        try:
            # Удаляем существующие данные
            await self.opensearch_service.delete_by_resource_id(
                resource.get("resource_id"), "vectorized_knowledge_resource"
            )

            query = resource['extra'].get("query")
            batch_size = resource['extra'].get("batch_size", 100)

            # Выполняем запрос и получаем все результаты
            all_records = await connection.fetch(query)
            total_records = len(all_records)

            logger.info(f"Found {total_records} records to process")

            # Обрабатываем записи батчами
            for i in range(0, total_records, batch_size):
                batch_data = all_records[i:i + batch_size]
                batch_number = (i // batch_size) + 1

                logger.info(f"Processing batch {batch_number} with {len(batch_data)} records")

                # Обрабатываем каждую запись в батче
                for record in batch_data:
                    record_dict = dict(record)
                    content_str = json.dumps(record_dict, cls=DecimalEncoder, ensure_ascii=False)
                    vector = self.vectorization_service.vectorize_data(content_str)

                    await self.opensearch_service.add_document(
                        resource.get("resource_id"),
                        "vectorized_knowledge_resource",
                        vector,
                        record_dict,
                        str(uuid.uuid4())
                    )

                logger.info(f"Completed batch {batch_number}")

        finally:
            await connection.close()


class MSSQLProcessingStrategy(ResourceProcessingStrategy):
    """Стратегия обработки MS SQL Server ресурсов"""

    async def process_resource(self, resource: Dict[str, Any]):
        logger.info(f"Processing MSSQL resource: {resource.get('resource_id')}")

        secret = self._get_secret(resource['extra'].get("database_secret_path"))
        connection_params = secret.get("connection_params")

        connection = pymssql.connect(
            server=connection_params.get("host"),
            user=connection_params.get("user"),
            password=connection_params.get("password"),
            database=connection_params.get("database"),
            port=connection_params.get("port")
        )

        try:
            # Удаляем существующие данные
            await self.opensearch_service.delete_by_resource_id(
                resource.get("resource_id"), "vectorized_knowledge_resource"
            )

            cursor = connection.cursor(as_dict=True)
            cursor.execute(resource['extra'].get("query"))
            batch_size = resource['extra'].get("batch_size", 100)

            while True:
                batch_data = cursor.fetchmany(batch_size)
                if not batch_data:
                    break

                # Обрабатываем каждую строку в батче
                for row in batch_data:
                    content_str = json.dumps(row, cls=DecimalEncoder, ensure_ascii=False)
                    vector = self.vectorization_service.vectorize_data(content_str)

                    await self.opensearch_service.add_document(
                        resource.get("resource_id"),
                        "vectorized_knowledge_resource",
                        vector,
                        row,
                        str(uuid.uuid4())
                    )

                logger.info(f"Processed batch of {len(batch_data)} records")

        finally:
            connection.close()


class GoogleDriveProcessingStrategy(ResourceProcessingStrategy):
    """Стратегия обработки Google Drive ресурсов"""

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self._drive_client = None
        self.supported_mime_types = {
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document': 'docx',
            'application/vnd.openxmlformats-officedocument.presentationml.presentation': 'pptx',
            'application/vnd.google-apps.document': 'docx',
        }

    def _get_drive_client(self):
        """Получение клиента Google Drive"""
        if not self._drive_client:
            secrets = self._get_secret('services/google_drive')
            credentials_info = json.loads(secrets.get("google_drive_credentials"))

            credentials = service_account.Credentials.from_service_account_info(
                credentials_info,
                scopes=['https://www.googleapis.com/auth/drive.readonly']
            )

            self._drive_client = googleapiclient.discovery.build(
                serviceName="drive", version="v3", credentials=credentials
            )

        return self._drive_client

    def _extract_folder_id(self, resource_path: str) -> str:
        """Извлечение ID папки из URL"""
        if 'drive.google.com' in resource_path:
            if 'folders/' in resource_path:
                return resource_path.split('folders/')[-1].split('?')[0]
            if 'id=' in resource_path:
                return resource_path.split('id=')[-1].split('&')[0]
        return resource_path

    def _get_file_type(self, filename: str) -> str:
        """Определение типа файла по расширению"""
        if filename.lower().endswith('.docx'):
            return 'docx'
        elif filename.lower().endswith(('.ppt', '.pptx')):
            return 'pptx'
        else:
            raise ValueError(f"Unsupported file type: {filename}")

    def _extract_docx_content(self, binary_data: bytes) -> str:
        """Извлечение текста из DOCX файла"""
        doc = Document(BytesIO(binary_data))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        return "\n".join(paragraphs)

    def _extract_pptx_content(self, binary_data: bytes) -> str:
        """Извлечение текста из PPTX файла"""
        presentation = Presentation(BytesIO(binary_data))
        text_runs = []
        for slide in presentation.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        if paragraph.text.strip():
                            text_runs.append(paragraph.text)
        return "\n".join(text_runs)

    def _extract_content(self, binary_data: bytes, file_type: str) -> str:
        """Извлечение контента из файла"""
        if file_type == 'docx':
            return self._extract_docx_content(binary_data)
        elif file_type == 'pptx':
            return self._extract_pptx_content(binary_data)
        return ""

    async def _process_file(self, file_id: str, file_name: str) -> Optional[str]:
        """Обработка одного файла"""
        try:
            drive_client = self._get_drive_client()
            request = drive_client.files().get_media(fileId=file_id)

            file_data = BytesIO()
            downloader = MediaIoBaseDownload(file_data, request)

            done = False
            while not done:
                _, done = downloader.next_chunk()

            binary_data = file_data.getvalue()
            file_type = self._get_file_type(file_name)
            content = self._extract_content(binary_data, file_type)

            logger.info(f"Successfully processed file: {file_name}")
            return content

        except Exception as e:
            logger.error(f"Failed to process file {file_name}: {e}")
            return None

    async def process_resource(self, resource: Dict[str, Any]):
        logger.info(f"Processing Google Drive resource: {resource.get('resource_id')}")

        try:
            # Удаляем существующие данные
            await self.opensearch_service.delete_by_resource_id(
                resource.get("resource_id"), "vectorized_knowledge_resource"
            )

            drive_client = self._get_drive_client()
            folder_id = self._extract_folder_id(resource.get("resource_path"))

            # Получаем список файлов в папке
            query = f"'{folder_id}' in parents and trashed = false"
            results = drive_client.files().list(
                q=query,
                fields="files(id, name, mimeType)"
            ).execute()

            files = results.get('files', [])
            logger.info(f"Found {len(files)} files in folder")

            for file_info in files:
                if file_info['mimeType'] in self.supported_mime_types:
                    content = await self._process_file(file_info['id'], file_info['name'])

                    if content and content.strip():
                        vector = self.vectorization_service.vectorize_data(content)

                        await self.opensearch_service.add_document(
                            resource.get("resource_id"),
                            "vectorized_knowledge_resource",
                            vector,
                            {
                                'filename': file_info['name'],
                                'content': content,
                                'mime_type': file_info['mimeType']
                            },
                            str(uuid.uuid4())
                        )

        except Exception as e:
            logger.error(f"Failed to process Google Drive resource: {e}")
            raise


class DynamoDBProcessingStrategy(ResourceProcessingStrategy):
    """Стратегия обработки DynamoDB ресурсов"""

    async def process_resource(self, resource: Dict[str, Any]):
        logger.info(f"Processing DynamoDB resource: {resource.get('resource_id')}")

        # Поскольку мы отходим от AWS, здесь можно добавить альтернативную реализацию
        # или просто логировать, что этот тип ресурсов не поддерживается
        logger.warning("DynamoDB processing not implemented - moving away from AWS services")


class ResourceProcessorFactory:
    """Фабрика для создания стратегий обработки ресурсов"""

    def __init__(self, vault_client: hvac.Client, vectorization_service: VectorizationService,
                 opensearch_service: OpenSearchService):
        self.vault_client = vault_client
        self.vectorization_service = vectorization_service
        self.opensearch_service = opensearch_service

        self._strategies = {
            'MYSQL': MySQLProcessingStrategy,
            'POSTGRESQL': PostgreSQLProcessingStrategy,
            'MSSQL': MSSQLProcessingStrategy,
            'GOOGLE_DRIVE': GoogleDriveProcessingStrategy,
            'DYNAMODB': DynamoDBProcessingStrategy,
        }

    def get_strategy(self, resource_type: str) -> ResourceProcessingStrategy:
        """Получение стратегии для типа ресурса"""
        strategy_class = self._strategies.get(resource_type)
        if not strategy_class:
            raise ValueError(f"Unsupported resource type: {resource_type}")

        return strategy_class(
            self.vault_client,
            self.vectorization_service,
            self.opensearch_service
        )


class UniversalVectorizationScheduler:
    """Универсальный планировщик векторизации с поддержкой всех типов ресурсов"""

    def __init__(self):
        self.scheduler = BlockingScheduler()
        self.vault_client = self._init_vault_client()
        self.mongo_client = AsyncMongoClient(
            MONGO_URL
        )

        # Инициализация сервисов
        self.vectorization_service = VectorizationService(self.vault_client)
        self.opensearch_service = OpenSearchService(self.vault_client)
        self.processor_factory = ResourceProcessorFactory(
            self.vault_client,
            self.vectorization_service,
            self.opensearch_service
        )

    def _init_vault_client(self) -> hvac.Client:
        """Инициализация клиента Vault"""
        try:
            client = hvac.Client(url=VAULT_URL, token=VAULT_TOKEN)
            # Настройте аутентификацию согласно вашей конфигурации
            logger.info("Vault client initialized successfully")
            return client
        except Exception as e:
            logger.error(f"Failed to initialize Vault client: {e}")
            raise

    def _get_mongo_client(self) -> AsyncMongoClient:
        """Получение клиента MongoDB"""
        if not self.mongo_client:
            secrets = self.vault_client.secrets.kv.v2.read_secret_version(
                path='database/mongodb'
            )['data']['data']
            connection_string = secrets.get('connection_string')
            self.mongo_client = AsyncMongoClient(connection_string)
        return self.mongo_client

    async def _get_all_resources(self) -> List[Dict[str, Any]]:
        """Получение всех ресурсов из MongoDB"""
        try:
            client = self._get_mongo_client()
            db = client[MONGO_DB_NAME]
            collection = db[MONGO_COLLECTION_NAME]

            resources = await collection.find().to_list()
            logger.info(f"Found {len(resources)} total resources")
            return resources

        except Exception as e:
            logger.error(f"Failed to get resources: {e}")
            return []

    async def _process_single_resource(self, resource: Dict[str, Any]):
        """Обработка одного ресурса"""
        try:
            resource_type = resource.get('type')
            resource_id = resource.get('resource_id')

            logger.info(f"Processing resource {resource_id} of type {resource_type}")

            strategy = self.processor_factory.get_strategy(resource_type)
            await strategy.process_resource(resource)

            logger.info(f"Successfully processed resource {resource_id}")

        except Exception as e:
            logger.error(f"Failed to process resource {resource.get('resource_id')}: {e}")

    async def _vectorize_all_resources(self):
        """Векторизация всех ресурсов"""
        try:
            logger.info("Starting universal vectorization process")

            resources = await self._get_all_resources()
            if not resources:
                logger.warning("No resources found")
                return

            # Группируем ресурсы по типам для логирования
            resource_counts = {}
            for resource in resources:
                resource_type = resource.get('type')
                resource_counts[resource_type] = resource_counts.get(resource_type, 0) + 1

            logger.info(f"Resource distribution: {resource_counts}")

            # Обрабатываем ресурсы параллельно (с ограничением)
            semaphore = asyncio.Semaphore(3)  # Максимум 3 параллельных обработки

            async def process_with_semaphore(resource):
                async with semaphore:
                    await self._process_single_resource(resource)

            tasks = [process_with_semaphore(resource) for resource in resources]
            await asyncio.gather(*tasks, return_exceptions=True)

            logger.info("Universal vectorization process completed")

        except Exception as e:
            logger.error(f"Vectorization process failed: {e}")
        finally:
            if self.mongo_client:
                await self.mongo_client.close()

    def vectorize_job(self):
        """Job функция для планировщика"""
        logger.info("Starting scheduled vectorization job")
        asyncio.run(self._vectorize_all_resources())
        logger.info("Scheduled vectorization job completed")

    def start(self):
        """Запуск планировщика"""
        logger.info(f"Starting universal vectorization scheduler with {SCHEDULER_INTERVAL_HOURS}h interval")

        # Основная задача
        self.scheduler.add_job(
            func=self.vectorize_job,
            trigger=IntervalTrigger(hours=SCHEDULER_INTERVAL_HOURS),
            id='universal_vectorize_job',
            name='Universal Vectorization Job',
            replace_existing=True,
            max_instances=1
        )

        # Первый запуск через 30 секунд
        self.scheduler.add_job(
            func=self.vectorize_job,
            trigger='date',
            run_date=datetime.datetime.now() + datetime.timedelta(seconds=30),
            id='initial_vectorize_job',
            name='Initial Universal Vectorization Job'
        )

        try:
            logger.info("Universal scheduler started successfully")
            self.scheduler.start()
        except KeyboardInterrupt:
            logger.info("Scheduler stopped by user")
            self.scheduler.shutdown()
        except Exception as e:
            logger.error(f"Scheduler error: {e}")
            self.scheduler.shutdown()


if __name__ == "__main__":
    scheduler = UniversalVectorizationScheduler()
    scheduler.start()

